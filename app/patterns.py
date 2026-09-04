"""Regex patterns for detecting Name / Email / Phone Number inside text or
uploaded document files, plus simple validators used by the contact form.

Run this file directly (``python -m app.patterns``) to execute the built-in
self-tests and confirm every pattern still matches/rejects the expected
sample values.
"""

import re
from pathlib import Path
from typing import Any

EMAIL_PATTERN = re.compile(
    r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
)

PHONE_PATTERN = re.compile(
    r"(?:\+92[\s-]?|0)3\d{2}[\s-]?\d{7}"
)

NAME_PATTERN = re.compile(
    r"\b[A-Z][a-z]+(?:[ \t]+[A-Z][a-z]+){1,2}\b"
)

NAME_FORM_PATTERN = re.compile(
    r"^[A-Za-z]+(?:[ \t.'-][A-Za-z]+)*$"
)

GENDER_PATTERN = re.compile(r"^(?:male|female)$", re.IGNORECASE)
SIGNATURE_PATTERN = re.compile(
    r"^data:image/png;base64,[A-Za-z0-9+/]+={0,2}$"
)


PLACEHOLDER_PATTERNS = {
    "name": re.compile(r"\{name\}", re.IGNORECASE),
    "email": re.compile(r"\[email\]", re.IGNORECASE),
    "contact": re.compile(r"\((?:contact|phone)\)", re.IGNORECASE),
}

PATTERNS = {
    "name": NAME_PATTERN,
    "email": EMAIL_PATTERN,
    "phone": PHONE_PATTERN,
}

FIELD_VALIDATORS = {
    "name": NAME_FORM_PATTERN,
    "email": EMAIL_PATTERN,
    "phone": PHONE_PATTERN,
    "gender": GENDER_PATTERN,
    "signature": SIGNATURE_PATTERN,
}

# Same regex sources the contact form uses — attached to CV fields by type.
CV_CONTACT_VALIDATORS = ("name", "email", "phone")
# Non-empty content (JS-safe; no Python-only flags). Used for skills, education, etc.
CV_DEFAULT_REGEX = r"^[\s\S]*\S[\s\S]*$"
CV_DEFAULT_PATTERN = re.compile(CV_DEFAULT_REGEX)


def fill_template(template: str, name: str, email: str, phone: str) -> str:
    """Replace {name}, [email], and (contact)/(phone) in template text."""
    text = PLACEHOLDER_PATTERNS["name"].sub(name, template)
    text = PLACEHOLDER_PATTERNS["email"].sub(email, text)
    text = PLACEHOLDER_PATTERNS["contact"].sub(phone, text)
    return text


def create_contact_document(
    target_path: Path,
    name: str,
    email: str,
    phone: str,
    gender: str | None = None,
    signature_path: Path | None = None,
) -> Path:
    """Create a Word document containing the submitted contact details."""
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_heading("CYBERIFY", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title.runs:
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(99, 102, 241)

    subtitle = doc.add_paragraph("Official Contact & Client Registration Form")
    subtitle.paragraph_format.space_after = Pt(18)
    for run in subtitle.runs:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(100, 116, 139)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Field Description"
    hdr_cells[1].text = "Details"
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(30, 41, 59)

    records = [
        ("Full Name", name),
        ("Email Address", email),
        ("Phone Number", phone),
    ]
    if gender:
        records.append(("Gender", gender.capitalize()))

    for label, val in records:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].text = val

    if signature_path and signature_path.exists():
        signature_heading = doc.add_heading("E-Signature", level=1)
        signature_heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        signature_heading.paragraph_format.space_after = Pt(0)
        signature_heading.paragraph_format.keep_with_next = True
        signature_paragraph = doc.add_paragraph()
        signature_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        signature_paragraph.paragraph_format.space_before = Pt(0)
        signature_paragraph.paragraph_format.space_after = Pt(0)
        signature_paragraph.add_run().add_picture(str(signature_path), height=Inches(1.0))

    target_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target_path))
    return target_path

def detect_patterns(text: str) -> dict:
    """Run every regex pattern against `text` and return the matches found."""
    results = {}
    for label, pattern in PATTERNS.items():
        matches = sorted(set(m.strip() for m in pattern.findall(text)))
        results[label] = matches
    return results


def compare_patterns(text: str, first: str, second: str) -> dict:
    """Compare the matches produced by two registered patterns."""
    first_pattern = PATTERNS.get(first)
    second_pattern = PATTERNS.get(second)
    if first_pattern is None or second_pattern is None:
        raise ValueError("Patterns must be registered name, email, or phone")

    first_matches = set(m.strip() for m in first_pattern.findall(text))
    second_matches = set(m.strip() for m in second_pattern.findall(text))
    return {
        "first": first,
        "second": second,
        "common": sorted(first_matches & second_matches),
        "only_in_first": sorted(first_matches - second_matches),
        "only_in_second": sorted(second_matches - first_matches),
    }


def validate_field(field: str, value: str) -> bool:
    """True if `value` fully matches the pattern registered for `field`."""
    pattern = FIELD_VALIDATORS.get(field)
    if pattern is None:
        raise ValueError(f"Unknown field: {field}")
    val = value.strip()
    if not val:
        return False
    return bool(pattern.fullmatch(val))


def validator_sources() -> dict[str, str]:
    """Regex strings for the browser (same patterns as FIELD_VALIDATORS)."""
    sources = {name: FIELD_VALIDATORS[name].pattern for name in CV_CONTACT_VALIDATORS}
    sources["content"] = CV_DEFAULT_REGEX
    return sources


def cv_field_regex(validator: str | None) -> str:
    if validator in FIELD_VALIDATORS:
        return FIELD_VALIDATORS[validator].pattern
    return CV_DEFAULT_REGEX


def infer_cv_validator(label: str, key: str = "", field_type: str = "") -> str | None:
    """Map a detected CV field onto name / email / phone validators when it is contact data."""
    kind = (field_type or "").lower()
    blob = f"{key} {label}".lower()
    if kind == "email" or "email" in blob:
        return "email"
    if kind in {"tel", "phone"} or re.search(r"\b(phone|mobile|whatsapp)\b", blob):
        return "phone"
    if kind == "name" or key in {"name", "full_name"} or re.search(
        r"\b(full[ _-]?name|candidate[ _-]?name)\b", blob
    ):
        return "name"
    return None


def validate_cv_fields(fields: list[dict[str, Any]]) -> list[dict[str, str]]:
    """Validate contact-like CV fields with the same regex as the contact form."""
    errors: list[dict[str, str]] = []
    for field in fields:
        label = str(field.get("label", "")).strip()
        key = str(field.get("key", "")).strip()
        field_type = str(field.get("type", "")).strip()
        validator = field.get("validator") or infer_cv_validator(label, key, field_type)
        value = str(field.get("value", ""))
        if validator in CV_CONTACT_VALIDATORS:
            ok = validate_field(validator, value)
        else:
            ok = bool(CV_DEFAULT_PATTERN.fullmatch(value.strip()))
        if not ok:
            errors.append({
                "label": label or key or "Field",
                "validator": validator or "content",
            })
    return errors



def extract_text_from_file(path: Path) -> str:
    """Best-effort plain-text extraction for the file types we support."""
    suffix = path.suffix.lower()

    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".docx":
        try:
            import docx  
        except ImportError as e:
            raise RuntimeError(
                "python-docx is required to read .docx files. "
                "Install it with: pip install python-docx"
            ) from e
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]
        table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(paragraphs + table_text)

    raise ValueError(f"Unsupported file type for pattern detection: {suffix}")


def _field_key(label: str, index: int) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
    return key or f"field_{index}"


def extract_cv_fields(text: str) -> list[dict[str, Any]]:
    """Turn common CV labels and section headings into editable form fields."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    fields: list[dict[str, Any]] = []
    seen: set[str] = set()
    section_names = {
        "profile", "summary", "objective", "experience", "work experience",
        "education", "skills", "projects", "certifications", "languages",
        "awards", "achievements", "interests", "references", "contact",
    }
    label_pattern = re.compile(r"^([A-Za-z][A-Za-z /&.-]{1,38})\s*:\s*(.*)$")

    def add_field(label: str, value: str, field_type: str = "text") -> None:
        clean_label = re.sub(r"\s+", " ", label).strip().strip("-|")
        clean_value = value.strip()
        if not clean_label or not clean_value:
            return
        key = _field_key(clean_label, len(fields))
        if key in seen:
            key = f"{key}_{len(fields) + 1}"
        seen.add(key)
        validator = infer_cv_validator(clean_label, key, field_type)
        regex = cv_field_regex(validator)
        fields.append({
            "key": key,
            "label": clean_label,
            "value": clean_value,
            "type": field_type,
            "validator": validator,
            "regex": regex,
        })

    for index, line in enumerate(lines):
        match = label_pattern.match(line)
        if match:
            label, value = match.groups()
            kind = "email" if label.lower() == "email" else "tel" if label.lower() in {"phone", "mobile", "contact"} else "text"
            add_field(label, value, kind)
            continue
        normalized = line.lower().rstrip(":")
        if normalized in section_names:
            value_lines: list[str] = []
            for following in lines[index + 1:]:
                if following.lower().rstrip(":") in section_names or label_pattern.match(following):
                    break
                value_lines.append(following)
            add_field(line.rstrip(":"), "\n".join(value_lines), "textarea")

    fallback_fields: list[dict[str, Any]] = []
    if not any(field.get("validator") == "name" for field in fields):
        names = NAME_PATTERN.findall("\n".join(lines[:4]))
        if names:
            fallback_fields.append({"label": "Full Name", "value": names[0], "type": "text"})
    if not any(field["type"] == "email" for field in fields):
        emails = EMAIL_PATTERN.findall(text)
        if emails:
            fallback_fields.append({"label": "Email", "value": emails[0], "type": "email"})
    if not any(field["type"] == "tel" for field in fields):
        phones = PHONE_PATTERN.findall(text)
        if phones:
            fallback_fields.append({"label": "Phone", "value": phones[0], "type": "tel"})

    if fallback_fields:
        original_fields = fields
        fields = []
        for fallback in fallback_fields:
            add_field(fallback["label"], fallback["value"], fallback["type"])
        fields.extend(original_fields)

    if not fields:
        add_field("CV Content", text.strip(), "textarea")
    return fields


def create_cv_document(target_path: Path, fields: list[dict[str, Any]], title: str = "CV") -> Path:
    """Create an editable CV DOCX from the fields submitted by the browser."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 64, 175)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for field in fields:
        row = table.add_row().cells
        row[0].text = str(field.get("label", "Field"))
        row[1].text = str(field.get("value", ""))
        row[0].paragraphs[0].runs[0].font.bold = True

    target_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target_path))
    return target_path



def _run_self_tests() -> None:
    sample_text = """
    Contact Person: Ali Raza
    Email me at ali.raza@example.com or hr.team@cyberify.io
    Call on 03001234567 or +92 300 1234567 for queries.
    Backup contact: Sara Ahmed, sara_ahmed99@gmail.com, 0345-9876543
    """

    results = detect_patterns(sample_text)
    print("Detected names :", results["name"])
    print("Detected emails:", results["email"])
    print("Detected phones:", results["phone"])

    assert "Ali Raza" in results["name"], "Name pattern failed"
    assert "Sara Ahmed" in results["name"], "Name pattern failed"
    assert "ali.raza@example.com" in results["email"], "Email pattern failed"
    assert "hr.team@cyberify.io" in results["email"], "Email pattern failed"
    assert "03001234567" in results["phone"], "Phone pattern failed"
    assert "0345-9876543" in results["phone"], "Phone pattern failed"

    # Field validators (used by the submit-form endpoint)
    assert validate_field("email", "student@cyberify.io") is True
    assert validate_field("email", "not-an-email") is False
    assert validate_field("phone", "03211234567") is True
    assert validate_field("phone", "12345") is False
    assert validate_field("name", "Ahmed Khan") is True
    assert validate_field("name", "Hamza") is True
    assert validate_field("name", "1234") is False

    print("\nAll pattern self-tests passed [OK]")


if __name__ == "__main__":
    _run_self_tests()
